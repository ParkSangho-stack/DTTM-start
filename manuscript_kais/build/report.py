"""수정내역서(.docx) 생성: changes/<판>.json -> DTTM_한국산학기술학회_수정내역서_<판>_<날짜>.docx

사용법: python3 build/report.py v1.0
"""
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
LATIN, HANGUL = "Arial", "맑은 고딕"
HEAD_FILL, DONE_FILL, TODAY_FILL = "1F3864", "E2EFDA", "FFF2CC"


def font(run, size=None, bold=None, color=None):
    run.font.name = LATIN
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), HANGUL)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def para(doc, text="", size=10, bold=False, color=None, align=None, after=4):
    p = doc.add_paragraph()
    if text:
        font(p.add_run(text), size, bold, color)
    p.paragraph_format.space_after = Pt(after)
    if align:
        p.alignment = align
    return p


def heading(doc, text):
    p = para(doc, text, 13, True, HEAD_FILL, after=6)
    p.paragraph_format.space_before = Pt(12)
    return p


def bullets(doc, items, size=10):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        font(p.add_run(it), size)
        p.paragraph_format.space_after = Pt(2)


def table(doc, header, rows, widths, size=9, fills=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        font(c.paragraphs[0].add_run(h), size, True, "FFFFFF")
        shade(c, HEAD_FILL)
    for r_i, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            lines = str(v).split("\n")
            p = cells[i].paragraphs[0]
            for k, line in enumerate(lines):
                if k:
                    p = cells[i].add_paragraph()
                font(p.add_run(line), size)
            if fills and fills[r_i]:
                shade(cells[i], fills[r_i])
    t.autofit = False
    for i, w in enumerate(widths):
        t.columns[i].width = Mm(w)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Mm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def main():
    ver = sys.argv[1]
    data = json.loads((ROOT / "changes" / f"{ver}.json").read_text())
    schedule = json.loads((ROOT / "build" / "schedule.json").read_text())
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(18)
    sec.top_margin = sec.bottom_margin = Mm(18)
    st = doc.styles["Normal"]
    st.font.name = LATIN
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), HANGUL)
    st.font.size = Pt(10)

    para(doc, "DTTM 투고 원고 수정내역서", 18, True, HEAD_FILL, WD_ALIGN_PARAGRAPH.CENTER, 2)
    para(doc, f"한국산학기술학회논문지 투고 준비 · {ver} ({data['date']})", 11, False, "595959",
         WD_ALIGN_PARAGRAPH.CENTER, 10)

    table(doc, ["항목", "내용"], [
        ["원고", data["manuscript"]],
        ["이번 판", f"{ver} ({data['date']}) — {data['topic']}"],
        ["기준 판", data["base"]],
        ["대상 학술지", "한국산학기술학회논문지 (Journal of the Korea Academia-Industrial cooperation Society)"],
        ["최종 목표", "2026-11-11 v2.0 최종본 동결 후 투고"],
    ], [35, 139])

    heading(doc, "1. 오늘의 수정 요약")
    bullets(doc, data["summary"])

    heading(doc, "2. 수정 내역 (수정 전 → 수정 후)")
    rows = [[str(i + 1), c["loc"], c["before"], c["after"], c["reason"]]
            for i, c in enumerate(data["changes"])]
    table(doc, ["No", "위치", "수정 전", "수정 후", "근거·이유"], rows, [9, 26, 48, 50, 41], 8.5)

    if data.get("metrics"):
        heading(doc, "3. 분량 현황")
        table(doc, ["지표", "값"], data["metrics"], [70, 104])

    heading(doc, "4. 저자 확인이 필요한 사항")
    bullets(doc, data["unverified"])

    heading(doc, "5. 바꾸지 않은 것")
    bullets(doc, data["unchanged"])

    heading(doc, "6. 다음 판 계획")
    bullets(doc, data["next"])

    heading(doc, "7. 전체 일정과 진행 상태")
    rows, fills = [], []
    for d, v, topic in schedule:
        status = "완료" if d < data["date"] else ("오늘" if d == data["date"] else "예정")
        rows.append([d, v, topic, status])
        fills.append(DONE_FILL if status == "완료" else (TODAY_FILL if status == "오늘" else None))
    table(doc, ["날짜", "판", "주제", "상태"], rows, [24, 15, 120, 15], 8.5, fills)

    sec_no = 8
    for extra in data.get("extra_tables", []):
        heading(doc, f'{sec_no}. {extra["title"]}')
        sec_no += 1
        if extra.get("note"):
            para(doc, extra["note"], 9, after=4)
        table(doc, extra["header"], extra["rows"], extra["widths"], 8)

    if data.get("sources"):
        heading(doc, f"{sec_no}. 확인한 자료")
        bullets(doc, data["sources"], 9)

    out = ROOT / f"DTTM_한국산학기술학회_수정내역서_{ver}_{data['date']}.docx"
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
